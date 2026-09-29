#!/usr/bin/env python3
"""Local GUI/CLI for extracting DJI video telemetry without touching the live backend."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import queue
import statistics
import sys
import tempfile
import threading
import traceback
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone, tzinfo
from pathlib import Path
from typing import Any, Callable, Iterable, Iterator, Mapping, Sequence


SERVICE_ROOT = Path(__file__).resolve().parent
if str(SERVICE_ROOT) not in sys.path:
    sys.path.insert(0, str(SERVICE_ROOT))
LOCAL_TCL_ROOT = SERVICE_ROOT / ".venv" / "tcl_runtime"
if (LOCAL_TCL_ROOT / "tcl8.6" / "init.tcl").is_file():
    os.environ.setdefault("TCL_LIBRARY", str(LOCAL_TCL_ROOT / "tcl8.6"))
if (LOCAL_TCL_ROOT / "tk8.6" / "tk.tcl").is_file():
    os.environ.setdefault("TK_LIBRARY", str(LOCAL_TCL_ROOT / "tk8.6"))

from backend_v1.dji_telemetry_extractor import (  # noqa: E402
    ExtractionCancelled as BackendExtractionCancelled,
    ExtractionError,
    ExtractionResult,
    TELEMETRY_VALUE_FIELDS,
    TelemetrySample,
    extract_dji_video,
)


APP_VERSION = "2.2.2"
LAUNCHER_LOG_PATH = SERVICE_ROOT / "dji_extractor_launcher.log"
AR_PROJECTION_MODE = "dji_srt_pose_terrain_intersection"
AR_REQUIRED_INPUTS = (
    "wgs84_position",
    "absolute_or_relative_altitude",
    "gimbal_yaw_pitch_roll",
    "video_sample_index",
    "video_time_ms",
    "estimated_intrinsics",
)
CSV_BASE_FIELDS = (
    "source_kind",
    "source_name",
    "video_sample_index",
    "frame_index",
    "cue_index",
    "start_time_ms",
    "end_time_ms",
    "recorded_at",
    *TELEMETRY_VALUE_FIELDS,
)


@dataclass(frozen=True)
class ExportOptions:
    raw_json: bool = True
    cesium_json: bool = True
    csv_file: bool = True
    xlsx_file: bool = True
    camera_json: bool = True
    ar_audit_json: bool = True


class ExtractionCancelled(ExtractionError):
    """A normal user-requested terminal state, not an extraction failure."""

    def __init__(
        self,
        message: str,
        *,
        result: ExtractionResult | None = None,
        generated: Sequence[Path] = (),
    ) -> None:
        super().__init__(message)
        self.result = result
        self.generated = tuple(generated)


def _parse_timezone(value: str) -> tuple[tzinfo, str]:
    text = value.strip()
    if not text or text.lower() == "local":
        local = datetime.now().astimezone().tzinfo or timezone.utc
        offset = datetime.now(local).strftime("%z")
        label = f"{offset[:3]}:{offset[3:]}" if offset else "UTC"
        return local, label
    if text.upper() in {"Z", "UTC", "+00", "+00:00"}:
        return timezone.utc, "+00:00"
    import re

    match = re.fullmatch(r"([+-])(\d{2})(?::?(\d{2}))?", text)
    if not match:
        raise ValueError("时区须填写 local、UTC 或 +08:00 形式")
    hours = int(match.group(2))
    minutes = int(match.group(3) or "0")
    if hours > 14 or minutes > 59 or (hours == 14 and minutes):
        raise ValueError("时区偏移超出有效范围")
    sign = 1 if match.group(1) == "+" else -1
    delta = timedelta(minutes=sign * (hours * 60 + minutes))
    return timezone(delta), f"{match.group(1)}{hours:02d}:{minutes:02d}"


def _atomic_bytes(destination: Path, content: bytes) -> Path:
    destination = destination.expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{destination.stem}-",
        suffix=f"{destination.suffix}.tmp",
        dir=destination.parent,
    )
    temporary = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "wb") as handle:
            handle.write(content)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, destination)
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise
    return destination


def _atomic_json(destination: Path, document: Any, *, compact: bool = False) -> Path:
    destination = destination.expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{destination.stem}-",
        suffix=f"{destination.suffix}.tmp",
        dir=destination.parent,
    )
    temporary = Path(temporary_name)
    encoder = json.JSONEncoder(
        ensure_ascii=False,
        separators=(",", ":") if compact else None,
        indent=None if compact else 2,
        allow_nan=False,
    )
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as handle:
            for chunk in encoder.iterencode(document):
                handle.write(chunk)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, destination)
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise
    return destination


def _sample_index(sample: TelemetrySample, fallback: int) -> int:
    explicit = getattr(sample, "video_sample_index", None)
    return int(explicit) if isinstance(explicit, int) and explicit >= 0 else fallback


def _primary_samples(result: ExtractionResult) -> tuple[TelemetrySample, ...]:
    """Choose one timed source for playback while the raw JSON keeps every source."""

    for source_kind in ("sidecar_srt", "embedded_srt", "ffprobe_tags"):
        selected = tuple(sample for sample in result.telemetry if sample.source_kind == source_kind)
        if selected:
            return selected
    return result.telemetry


def _recorded_epoch_ms(recorded_at: str | None, source_timezone: tzinfo) -> int | None:
    if not recorded_at:
        return None
    try:
        parsed = datetime.fromisoformat(recorded_at.replace("Z", "+00:00"))
    except ValueError:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=source_timezone)
    return round(parsed.timestamp() * 1000)


def _timeline_anchor_ms(samples: Sequence[TelemetrySample], source_timezone: tzinfo) -> int | None:
    for sample in samples:
        epoch = _recorded_epoch_ms(sample.recorded_at, source_timezone)
        if epoch is not None and sample.start_time_ms is not None:
            return epoch - sample.start_time_ms
    return None


def _focal_35mm(sample: TelemetrySample) -> float | None:
    value = getattr(sample, "focal_length_35mm_equivalent_mm", None)
    if value is None:
        value = getattr(sample, "focal_length_mm", None)
    return float(value) if isinstance(value, (int, float)) and math.isfinite(value) and value > 0 else None


def _camera_values(sample: TelemetrySample, width: int | None, height: int | None) -> dict[str, Any]:
    focal = _focal_35mm(sample)
    zoom = sample.digital_zoom_ratio
    result: dict[str, Any] = {}
    if focal is not None:
        result["focal_length_35mm_equivalent_mm"] = focal
    if zoom is not None and math.isfinite(zoom) and zoom > 0:
        result["digital_zoom_ratio"] = zoom
    effective = focal * (zoom if zoom is not None and zoom > 0 else 1.0) if focal is not None else None
    if effective is None or not width or not height:
        return result
    fx = width * effective / 36.0
    fy = fx
    result["estimated_intrinsics"] = {
        "fx_px": round(fx, 6),
        "fy_px": round(fy, 6),
        "cx_px": round(width / 2.0, 6),
        "cy_px": round(height / 2.0, 6),
        "horizontal_fov_degrees": round(math.degrees(2 * math.atan(width / (2 * fx))), 8),
        "vertical_fov_degrees": round(math.degrees(2 * math.atan(height / (2 * fy))), 8),
        "method": "35mm-equivalent focal length; centered principal point; square pixels",
        "distortion_coefficients": None,
        "quality": "estimated_not_calibrated",
    }
    return result


def _finite_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def _projection_input_audit(
    sample: TelemetrySample,
    camera: dict[str, Any],
) -> dict[str, bool]:
    """Audit only explicit source values; fallback export values must not imply readiness."""

    longitude = sample.longitude
    latitude = sample.latitude
    estimated = camera.get("estimated_intrinsics")
    has_intrinsics = isinstance(estimated, dict) and all(
        _finite_number(estimated.get(key))
        for key in ("fx_px", "fy_px", "cx_px", "cy_px", "horizontal_fov_degrees", "vertical_fov_degrees")
    )
    if has_intrinsics:
        has_intrinsics = (
            estimated["fx_px"] > 0
            and estimated["fy_px"] > 0
            and 0 < estimated["horizontal_fov_degrees"] < 180
            and 0 < estimated["vertical_fov_degrees"] < 180
            and estimated.get("quality") == "estimated_not_calibrated"
        )
    return {
        "wgs84_position": (
            _finite_number(longitude)
            and _finite_number(latitude)
            and -180 <= longitude <= 180
            and -90 <= latitude <= 90
        ),
        "absolute_or_relative_altitude": any(
            _finite_number(value)
            for value in (
                sample.absolute_altitude_m,
                sample.gps_altitude_m,
                sample.relative_altitude_m,
            )
        ),
        "gimbal_yaw_pitch_roll": all(
            _finite_number(value)
            for value in (
                sample.gimbal_yaw_deg,
                sample.gimbal_pitch_deg,
                sample.gimbal_roll_deg,
            )
        ),
        # These two checks intentionally inspect the original SRT values.  The
        # exporter still writes compatibility fallbacks, but a fallback never
        # makes an AR frame projection-ready.
        "video_sample_index": (
            isinstance(sample.video_sample_index, int)
            and not isinstance(sample.video_sample_index, bool)
            and sample.video_sample_index >= 0
        ),
        "video_time_ms": _finite_number(sample.start_time_ms) and sample.start_time_ms >= 0,
        "estimated_intrinsics": has_intrinsics,
    }


def _projection_audits(
    samples: Sequence[TelemetrySample],
    width: int | None,
    height: int | None,
) -> list[dict[str, bool]]:
    return [_projection_input_audit(sample, _camera_values(sample, width, height)) for sample in samples]


def _ar_projection_contract(audits: Sequence[dict[str, bool]]) -> dict[str, Any]:
    total = len(audits)
    coverage = {
        name: sum(bool(audit.get(name)) for audit in audits)
        for name in AR_REQUIRED_INPUTS
    }
    ready_points = sum(all(audit.get(name, False) for name in AR_REQUIRED_INPUTS) for audit in audits)
    missing_counts = {name: total - coverage[name] for name in AR_REQUIRED_INPUTS}
    ready = total > 0 and ready_points == total
    status = "ready_estimated" if ready else "partial" if ready_points else "not_ready"
    return {
        "mode": AR_PROJECTION_MODE,
        "readiness": {
            "ready": ready,
            "status": status,
            "quality": "estimated_not_survey_grade",
            "source_sample_count": total,
            "ready_point_count": ready_points,
            "coverage": coverage,
            "missing_input_counts": missing_counts,
            "required_inputs": list(AR_REQUIRED_INPUTS),
        },
        "runtime_requirements": {
            "terrain_intersection": True,
            "terrain_height_reference": "Cesium terrain or configured ground altitude",
            "video_binding": "video_sample_index + video_time_ms (SRT PTS)",
            "body_attitude_required": False,
        },
        "measured_outputs": {
            "ground_corners_available": False,
            "lens_distortion_coefficients": False,
        },
        "limitations": [
            "camera intrinsics are estimated from 35mm-equivalent focal length",
            "DJI SRT gimbal attitude is used directly; body attitude is not fabricated",
            "ground corners are calculated at runtime by terrain intersection, not measured",
        ],
    }


def _ground_altitude(samples: Sequence[TelemetrySample]) -> float | None:
    values = [
        sample.absolute_altitude_m - sample.relative_altitude_m
        for sample in samples
        if sample.absolute_altitude_m is not None and sample.relative_altitude_m is not None
        and math.isfinite(sample.absolute_altitude_m) and math.isfinite(sample.relative_altitude_m)
    ]
    return statistics.median(values) if values else None


def build_cesium_trajectory(
    result: ExtractionResult,
    *,
    source_timezone: tzinfo,
    timezone_label: str,
) -> dict[str, Any]:
    """Build a compact trajectory while keeping measured and estimated fields distinct."""

    samples = _primary_samples(result)
    anchor_ms = _timeline_anchor_ms(samples, source_timezone)
    width = result.video.width
    height = result.video.height
    projection_audits = _projection_audits(samples, width, height)
    points: list[dict[str, Any]] = []
    skipped_without_position = 0
    for fallback_index, sample in enumerate(samples):
        if sample.latitude is None or sample.longitude is None:
            skipped_without_position += 1
            continue
        absolute_altitude = sample.absolute_altitude_m
        if absolute_altitude is None:
            absolute_altitude = sample.gps_altitude_m
        altitude = absolute_altitude if absolute_altitude is not None else sample.relative_altitude_m
        if altitude is None:
            skipped_without_position += 1
            continue
        video_time_ms = sample.start_time_ms if sample.start_time_ms is not None else 0
        timestamp_ms = anchor_ms + video_time_ms if anchor_ms is not None else None
        camera = _camera_values(sample, width, height)
        projection_audit = projection_audits[fallback_index]
        projection_missing = [name for name in AR_REQUIRED_INPUTS if not projection_audit[name]]
        point: dict[str, Any] = {
            "video_sample_index": _sample_index(sample, fallback_index),
            "frame_index": sample.frame_index,
            "video_time_ms": video_time_ms,
            "recorded_at": sample.recorded_at,
            "position": {
                "longitude": sample.longitude,
                "latitude": sample.latitude,
                "altitude": altitude,
                "relative_altitude": sample.relative_altitude_m,
                "altitude_source": (
                    "DJI SRT abs_alt" if sample.absolute_altitude_m is not None
                    else "DJI SRT GPS altitude" if sample.gps_altitude_m is not None
                    else "DJI SRT rel_alt"
                ),
            },
            "gimbal_attitude": {
                key: value
                for key, value in {
                    "yaw": sample.gimbal_yaw_deg,
                    "pitch": sample.gimbal_pitch_deg,
                    "roll": sample.gimbal_roll_deg,
                }.items()
                if value is not None
            },
            "camera": camera,
            "ar_projection_ready": not projection_missing,
            "ar_projection_missing_inputs": projection_missing or None,
            "exposure": {
                key: value
                for key, value in {
                    "iso": sample.iso,
                    "shutter_seconds": sample.shutter_seconds,
                    "aperture_f_number": sample.aperture_f_number,
                    "exposure_value_ev": sample.exposure_value_ev,
                    "color_mode": sample.color_mode,
                    "ae_metering_mode": sample.ae_metering_mode,
                }.items()
                if value is not None
            },
        }
        if timestamp_ms is not None:
            point["source_timestamp_ms"] = timestamp_ms
            point["video_frame_timestamp_ms"] = timestamp_ms
            point["video_frame_time_delta_ms"] = 0
        else:
            point["timestamp"] = video_time_ms / 1000.0
        attitude = {
            key: value
            for key, value in {
                "heading": sample.aircraft_yaw_deg,
                "pitch": sample.aircraft_pitch_deg,
                "roll": sample.aircraft_roll_deg,
            }.items()
            if value is not None
        }
        if attitude:
            point["attitude"] = attitude
        points.append(point)

    ground_altitude = _ground_altitude(samples)
    focal_values = [_focal_35mm(sample) for sample in samples]
    focal_values = [value for value in focal_values if value is not None]
    zoom_values = [
        sample.digital_zoom_ratio for sample in samples
        if sample.digital_zoom_ratio is not None and sample.digital_zoom_ratio > 0
    ]
    camera_summary: dict[str, Any] = {
        "image_width_px": width,
        "image_height_px": height,
        "focal_length_35mm_equivalent_mm_range": (
            [min(focal_values), max(focal_values)] if focal_values else None
        ),
        "digital_zoom_ratio_range": [min(zoom_values), max(zoom_values)] if zoom_values else None,
        "intrinsics_status": "estimated_from_35mm_equivalent" if focal_values and width and height else "unavailable",
        "lens_distortion_status": "not_present_in_DJI_SRT",
    }
    document: dict[str, Any] = {
        "schema_version": "cesium-dji-trajectory/v1",
        "source_crs": "EPSG:4326",
        "flight_info": {
            "generator": f"DJI 航拍视频信息提取器 {APP_VERSION}",
            "video_file": result.source_file,
            "video": result.video.to_dict(),
            "timezone": timezone_label,
            "timeline": (
                "exact SRT PTS anchored to first DJI wall-clock sample"
                if anchor_ms is not None else "relative SRT PTS; absolute clock unavailable"
            ),
            "altitude_reference": "absolute" if any(
                sample.absolute_altitude_m is not None or sample.gps_altitude_m is not None
                for sample in samples
            ) else "relative",
            "ground_altitude_m": ground_altitude,
            "camera_model": camera_summary,
            "body_attitude_status": (
                "available" if any(sample.aircraft_yaw_deg is not None for sample in samples)
                else "not_present_in_DJI_SRT"
            ),
            "ground_footprint_status": "not_extracted; requires terrain intersection or surveyed control points",
            "skipped_samples_without_complete_position": skipped_without_position,
            "ar_projection": _ar_projection_contract(projection_audits),
        },
        "trajectory_points": points,
    }
    return _drop_none(document)


def _source_video_identity(video_path: Path, result: ExtractionResult) -> dict[str, Any]:
    resolved = video_path.expanduser().resolve()
    size_bytes = resolved.stat().st_size if resolved.is_file() else result.video.size_bytes
    duration_ms = result.video.duration_ms
    if duration_ms is None and result.video.frame_count and result.video.frame_rate_fps:
        duration_ms = round(result.video.frame_count / result.video.frame_rate_fps * 1000)
    identity: dict[str, Any] = {
        "file_name": resolved.name or Path(result.source_file).name,
        "size_bytes": int(size_bytes) if isinstance(size_bytes, int) and size_bytes > 0 else None,
        "duration_ms": int(duration_ms) if isinstance(duration_ms, int) and duration_ms > 0 else None,
        "rotation_degrees": result.video.rotation_deg,
        "width_px": result.video.width,
        "height_px": result.video.height,
        "frame_rate_fps": result.video.frame_rate_fps,
        "frame_count": result.video.frame_count,
    }
    if resolved.is_file() and isinstance(size_bytes, int) and size_bytes > 0:
        block_size = 1024 * 1024
        offsets = sorted({0, max(0, size_bytes // 2 - block_size // 2), max(0, size_bytes - block_size)})
        digest = hashlib.sha256()
        digest.update(b"dji-video-sampled-sha256-v1\0")
        digest.update(str(size_bytes).encode("ascii"))
        with resolved.open("rb") as handle:
            for offset in offsets:
                handle.seek(offset)
                payload = handle.read(min(block_size, size_bytes - offset))
                digest.update(offset.to_bytes(8, "big", signed=False))
                digest.update(len(payload).to_bytes(8, "big", signed=False))
                digest.update(payload)
        identity.update({
            "content_fingerprint_sha256": digest.hexdigest(),
            "content_fingerprint_scheme": "sha256(size + first/middle/last 1MiB)-v1",
        })
    return identity


def _camera_segments(result: ExtractionResult) -> list[dict[str, Any]]:
    samples = _primary_samples(result)
    width = result.video.width
    height = result.video.height
    segments: list[dict[str, Any]] = []
    current_key: str | None = None
    current: dict[str, Any] | None = None
    for fallback_index, sample in enumerate(samples):
        camera = _camera_values(sample, width, height)
        key = json.dumps(camera, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        sample_index = _sample_index(sample, fallback_index)
        video_time_ms = sample.start_time_ms if sample.start_time_ms is not None else 0
        if key != current_key:
            current = {
                "start_video_sample_index": sample_index,
                "end_video_sample_index": sample_index,
                "start_video_time_ms": video_time_ms,
                "end_video_time_ms": video_time_ms,
                "camera": camera,
            }
            segments.append(current)
            current_key = key
        elif current is not None:
            current["end_video_sample_index"] = sample_index
            current["end_video_time_ms"] = video_time_ms
    return segments


def build_camera_lens_profile(
    result: ExtractionResult,
    *,
    source_video: Mapping[str, Any],
) -> dict[str, Any]:
    """Describe measured lens metadata separately from estimated calibration."""

    samples = _primary_samples(result)
    focal_count = sum(_focal_35mm(sample) is not None for sample in samples)
    zoom_count = sum(
        isinstance(sample.digital_zoom_ratio, (int, float))
        and math.isfinite(sample.digital_zoom_ratio)
        and sample.digital_zoom_ratio > 0
        for sample in samples
    )
    profile = {
        "schema_version": "dji-video-camera-lens/v2",
        "generator": f"DJI 航拍视频 AR 数据提取器 {APP_VERSION}",
        "source_video": dict(source_video),
        "image_space": "browser_display",
        "image_width_px": result.video.width,
        "image_height_px": result.video.height,
        "rotation_degrees": result.video.rotation_deg,
        "sample_count": len(samples),
        "measured_coverage": {
            "focal_length_35mm_equivalent": focal_count,
            "digital_zoom_ratio": zoom_count,
        },
        "intrinsics": {
            "status": "estimated_from_35mm_equivalent" if focal_count and result.video.width and result.video.height else "unavailable",
            "per_sample_segments": _camera_segments(result),
            "principal_point": "image_center_estimate",
            "pixel_aspect_ratio": "1.0_estimate",
        },
        "lens_distortion": {
            "status": "not_present_in_DJI_SRT_or_container_metadata",
            "coefficients": None,
            "recovery_method": "SfM bundle adjustment in optional COLMAP reconstruction",
            "must_not_assume_zero_for_exact_projection": True,
        },
        "calibration_quality": "estimated_not_calibrated",
        "exact_projection_requirements": [
            "calibrated/refined intrinsics and lens distortion",
            "per-frame camera pose",
            "first-visible-surface depth or a surveyed 3D surface",
            "WGS84 ellipsoid-height georeferencing using RTK/PPK/GCP or confirmed altitude conversion",
        ],
    }
    return _drop_none(profile)


def _projection_sample_records(result: ExtractionResult) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    width = result.video.width
    height = result.video.height
    for fallback_index, sample in enumerate(_primary_samples(result)):
        camera = _camera_values(sample, width, height)
        record = {
            "video_sample_index": _sample_index(sample, fallback_index),
            "video_time_ms": sample.start_time_ms,
            "longitude": sample.longitude,
            "latitude": sample.latitude,
            "absolute_altitude_m": sample.absolute_altitude_m,
            "gps_altitude_m": sample.gps_altitude_m,
            "relative_altitude_m": sample.relative_altitude_m,
            "gimbal_yaw_deg": sample.gimbal_yaw_deg,
            "gimbal_pitch_deg": sample.gimbal_pitch_deg,
            "gimbal_roll_deg": sample.gimbal_roll_deg,
            "aircraft_yaw_deg": sample.aircraft_yaw_deg,
            "aircraft_pitch_deg": sample.aircraft_pitch_deg,
            "aircraft_roll_deg": sample.aircraft_roll_deg,
            "focal_length_35mm_equivalent_mm": _focal_35mm(sample),
            "digital_zoom_ratio": sample.digital_zoom_ratio,
            "estimated_intrinsics": camera.get("estimated_intrinsics"),
        }
        records.append(_drop_none(record))
    return records


def build_ar_preparation_document(
    result: ExtractionResult,
    *,
    source_video: Mapping[str, Any],
    projection_options: Any | None = None,
) -> dict[str, Any]:
    try:
        from projection_reconstruction import ProjectionOptions, projection_capabilities

        selected = projection_options if isinstance(projection_options, ProjectionOptions) else ProjectionOptions()
        capabilities = projection_capabilities(selected)
    except BaseException as exc:
        capabilities = {
            "ready": False,
            "missing": [f"三维重建模块不可用：{exc}"],
        }
    samples = _primary_samples(result)
    coverage = {
        "wgs84_position": sum(sample.longitude is not None and sample.latitude is not None for sample in samples),
        "absolute_or_gps_altitude": sum(
            sample.absolute_altitude_m is not None or sample.gps_altitude_m is not None for sample in samples
        ),
        "relative_altitude": sum(sample.relative_altitude_m is not None for sample in samples),
        "gimbal_yaw_pitch_roll": sum(all(value is not None for value in (
            sample.gimbal_yaw_deg, sample.gimbal_pitch_deg, sample.gimbal_roll_deg,
        )) for sample in samples),
        "body_yaw_pitch_roll": sum(all(value is not None for value in (
            sample.aircraft_yaw_deg, sample.aircraft_pitch_deg, sample.aircraft_roll_deg,
        )) for sample in samples),
        "exact_pts": sum(sample.start_time_ms is not None for sample in samples),
        "estimated_intrinsics": sum(bool(_camera_values(sample, result.video.width, result.video.height).get("estimated_intrinsics")) for sample in samples),
    }
    return _drop_none({
        "schema_version": "dji-video-ar-preparation/v2",
        "generator": f"DJI 航拍视频 AR 数据提取器 {APP_VERSION}",
        "source_video": dict(source_video),
        "source_sample_count": len(samples),
        "directly_extracted": coverage,
        "reconstruction_capabilities": capabilities,
        "required_projection_mesh_contract": {
            "schema": "cesium-video-projection-mesh/v1",
            "source_crs": "EPSG:4979",
            "height_reference": "WGS84_ELLIPSOID",
            "image_space": "browser_display",
            "visibility": "first_surface_hit",
            "vertex": "[u,v,longitude,latitude,ellipsoidHeight]",
            "binding": "video_sample_index + video_time_ms",
            "maximum_pts_gap_ms": 100,
        },
        "computed_not_metadata": {
            "calibrated_intrinsics_and_distortion": "optional SfM bundle adjustment",
            "visible_surface_depth": "optional MVS or external surveyed 3D surface",
            "per_frame_wgs84_mesh": "only published after full quality validation",
        },
        "safety": {
            "fabricate_ground_corners": False,
            "flatten_buildings_to_ground": False,
            "publish_partial_mesh": False,
            "unknown_vertical_datum_is_epsg4979": False,
        },
    })


def _drop_none(value: Any) -> Any:
    if isinstance(value, dict):
        return {key: _drop_none(item) for key, item in value.items() if item is not None}
    if isinstance(value, list):
        return [_drop_none(item) for item in value]
    return value


def _csv_layout(
    result: ExtractionResult,
) -> tuple[tuple[TelemetrySample, ...], list[str], tuple[str, ...]]:
    samples = _primary_samples(result)
    extra_names = tuple(sorted({key for sample in samples for key in sample.extra_fields}))
    headers = list(CSV_BASE_FIELDS) + [f"raw_{name}" for name in extra_names] + ["raw_text"]
    return samples, headers, extra_names


def _iter_csv_rows(
    samples: Sequence[TelemetrySample],
    extra_names: Sequence[str],
) -> Iterator[dict[str, Any]]:
    for fallback_index, sample in enumerate(samples):
        row = {name: getattr(sample, name, None) for name in CSV_BASE_FIELDS}
        row["video_sample_index"] = _sample_index(sample, fallback_index)
        row.update({f"raw_{name}": sample.extra_fields.get(name) for name in extra_names})
        row["raw_text"] = sample.raw_text
        yield row


def _csv_rows(result: ExtractionResult) -> tuple[list[str], list[dict[str, Any]]]:
    """Compatibility helper for callers/tests that explicitly need materialized rows."""

    samples, headers, extra_names = _csv_layout(result)
    return headers, list(_iter_csv_rows(samples, extra_names))


def write_csv(result: ExtractionResult, destination: Path) -> Path:
    samples, headers, extra_names = _csv_layout(result)
    destination = destination.expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(prefix=f".{destination.stem}-", suffix=".csv.tmp", dir=destination.parent)
    temporary = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(_iter_csv_rows(samples, extra_names))
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, destination)
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise
    return destination


def write_xlsx(result: ExtractionResult, destination: Path) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("导出 XLSX 需要 openpyxl，请重新执行 pip install -r requirements.txt") from exc
    samples, headers, extra_names = _csv_layout(result)
    workbook = Workbook(write_only=True)
    destination = destination.expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(f".{destination.stem}-{os.getpid()}.tmp.xlsx")
    try:
        sheet = workbook.create_sheet()
        sheet.title = "DJI逐帧遥测"
        last_column = get_column_letter(len(headers))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{last_column}{len(samples) + 1}"
        for index, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(36, max(12, len(header) + 2))

        fill = PatternFill("solid", fgColor="1F4E78")
        font = Font(color="FFFFFF", bold=True)
        header_cells = []
        for header in headers:
            cell = WriteOnlyCell(sheet, value=header)
            cell.font = font
            cell.fill = fill
            header_cells.append(cell)
        sheet.append(header_cells)
        for row in _iter_csv_rows(samples, extra_names):
            sheet.append([row.get(header) for header in headers])
        workbook.save(temporary)
        os.replace(temporary, destination)
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()
    return destination


def build_report(result: ExtractionResult, generated: Iterable[Path], timezone_label: str) -> str:
    generated_paths = list(generated)
    cesium_path = next((path for path in generated_paths if "Cesium轨迹" in path.name), None)
    camera_path = next((path for path in generated_paths if "相机内参与镜头信息" in path.name), None)
    ar_audit_path = next((path for path in generated_paths if "AR投影数据准备" in path.name), None)
    mesh_path = next((path for path in generated_paths if "三维投影网格_请选择此文件" in path.name), None)
    quality_path = next((path for path in generated_paths if "AR投影质量报告" in path.name), None)
    reconstruction_status = "未请求"
    if quality_path and quality_path.is_file():
        try:
            quality_document = json.loads(quality_path.read_text(encoding="utf-8"))
            reconstruction_status = str(quality_document.get("status") or "未知")
            if quality_document.get("message"):
                reconstruction_status += f"（{quality_document['message']}）"
        except (OSError, json.JSONDecodeError):
            reconstruction_status = "质量报告无法读取"
    samples = _primary_samples(result)
    ar_contract = _ar_projection_contract(
        _projection_audits(samples, result.video.width, result.video.height)
    )
    ar_readiness = ar_contract["readiness"]
    ar_status = (
        "已就绪（DJI SRT 姿态 + Cesium 地形求交，估算）"
        if ar_readiness["ready"]
        else f"未完全就绪（{ar_readiness['ready_point_count']}/{ar_readiness['source_sample_count']} 帧具备完整输入）"
    )
    counts = {
        name: sum(getattr(sample, name, None) is not None for sample in samples)
        for name in TELEMETRY_VALUE_FIELDS
    }
    video = result.video
    lines = [
        "DJI 航拍视频信息提取报告",
        f"程序版本: {APP_VERSION}",
        f"视频: {result.source_file}",
        f"视频参数: {video.width or '?'}×{video.height or '?'} / {video.frame_rate_fps or '?'} FPS / {video.frame_count or '?'} 帧",
        f"视频探测方式: {video.backend}",
        f"用于逐帧导出的遥测记录: {len(samples)} 条",
        f"原始来源记录总数: {len(result.telemetry)} 条",
        f"输入墙钟时区: {timezone_label}",
        f"GPS完整记录: {min(counts.get('latitude', 0), counts.get('longitude', 0))}",
        f"绝对/相对高度: {counts.get('absolute_altitude_m', 0)} / {counts.get('relative_altitude_m', 0)}",
        f"云台姿态: {min(counts.get('gimbal_yaw_deg', 0), counts.get('gimbal_pitch_deg', 0), counts.get('gimbal_roll_deg', 0))}",
        f"机身姿态: {min(counts.get('aircraft_yaw_deg', 0), counts.get('aircraft_pitch_deg', 0), counts.get('aircraft_roll_deg', 0))}",
        f"等效焦距/数字变焦: {counts.get('focal_length_35mm_equivalent_mm', counts.get('focal_length_mm', 0))} / {counts.get('digital_zoom_ratio', 0)}",
        f"AR 投影就绪状态: {ar_status}",
        f"真实表面 SfM/MVS 重建: {reconstruction_status}",
        "",
        "重要说明:",
        "- FrameCnt 原值保留，但视频绑定使用连续 video_sample_index 和 SRT PTS。",
        "- gb_yaw/gb_pitch/gb_roll 是云台姿态，不会伪装成机身姿态。",
        "- fx/fy/FOV 是根据35mm等效焦距估算；主点和畸变系数并未出现在当前 SRT 中。",
        "- 就绪时，Cesium 可使用逐帧 GPS/高度/云台三轴/PTS/估算内参进行“SRT 姿态 + 地形求交（估算）”。",
        "- 本工具不伪造地面四角；真实 WGS84 四角仍需相机标定、地形求交成果或地面控制点。",
        "- 镜头畸变、逐像素深度和建筑表面坐标不是 MP4/SRT 元数据；三维模式通过 SfM/MVS 计算并经过质量校验后才发布。",
        "- 垂直基准未确认时不会把 DJI 高度冒充 EPSG:4979 椭球高，也不会生成空的三维网格。",
        "",
        "Cesium 加载方法:",
        f"- 轨迹 JSON：{cesium_path if cesium_path else '本次未生成'}",
        f"- 本地视频：{result.source_file}",
        f"- 真实表面三维网格：{mesh_path if mesh_path else '本次未生成；可查看 AR 投影质量报告'}",
        "- 逐帧 WGS84 地面四角：本次未生成，请留空；三维模式请选择网格文件，不要选择“全部元数据”文件。",
        f"- 投影模式：{ar_status}",
        f"- 相机/镜头信息：{camera_path if camera_path else '本次未生成'}",
        f"- AR 数据准备审计：{ar_audit_path if ar_audit_path else '本次未生成'}",
        "",
        "生成文件:",
        *[f"- {path}" for path in generated_paths],
    ]
    if result.warnings:
        lines.extend(["", "警告:", *[f"- {warning}" for warning in result.warnings]])
    return "\n".join(lines) + "\n"


def run_extraction(
    video_path: Path,
    *,
    srt_path: Path | None,
    output_directory: Path,
    source_timezone_text: str,
    options: ExportOptions,
    ffprobe_bin: str = "ffprobe",
    ffmpeg_bin: str = "ffmpeg",
    projection_options: Any | None = None,
    progress_callback: Callable[[str, str, float | None], None] | None = None,
    cancel_event: threading.Event | None = None,
) -> tuple[ExtractionResult, list[Path]]:
    def check_cancelled() -> None:
        if cancel_event is not None and cancel_event.is_set():
            raise ExtractionCancelled("用户已取消提取")

    check_cancelled()
    if progress_callback:
        progress_callback("telemetry", "正在读取视频头和 DJI 逐帧遥测……", None)
    source_timezone, timezone_label = _parse_timezone(source_timezone_text)
    output_directory = output_directory.expanduser().resolve()
    output_directory.mkdir(parents=True, exist_ok=True)
    preflight_generated: list[Path] = []
    try:
        if srt_path is not None:
            from backend_v1.dji_srt_exporter import prepare_dji_srt

            preparation = prepare_dji_srt(
                video_path,
                output_directory,
                explicit_sidecar=srt_path,
                ffprobe_bin=ffprobe_bin,
                ffmpeg_bin=ffmpeg_bin,
                cancel_event=cancel_event,
            )
            if preparation.report_path:
                preflight_generated.append(Path(preparation.report_path))
            if not preparation.ar_ready or not preparation.output_srt:
                report_hint = (
                    f"；检查报告：{preparation.report_path}"
                    if preparation.report_path
                    else ""
                )
                raise ExtractionError(preparation.message + report_hint)
            srt_path = Path(preparation.output_srt)
        result = extract_dji_video(
            video_path,
            srt_path=srt_path,
            ffprobe_bin=ffprobe_bin,
            ffmpeg_bin=ffmpeg_bin,
            cancel_event=cancel_event,
        )
    except BackendExtractionCancelled as exc:
        raise ExtractionCancelled(str(exc)) from exc
    check_cancelled()
    if not _primary_samples(result):
        raise ExtractionError("没有提取到逐帧遥测；请确认已选择 DJI 同名 SRT 或包含遥测的字幕流")
    prefix = video_path.stem
    generated: list[Path] = list(preflight_generated)
    projection_success_committed = False
    source_video = _source_video_identity(video_path, result)
    check_cancelled()
    if options.raw_json:
        generated.append(_atomic_json(output_directory / f"{prefix}_全部元数据_仅供审计.json", result.to_dict(), compact=True))
        check_cancelled()
    if options.cesium_json:
        generated.append(_atomic_json(
            output_directory / f"{prefix}_Cesium轨迹_请选择此文件.json",
            build_cesium_trajectory(result, source_timezone=source_timezone, timezone_label=timezone_label),
            compact=True,
        ))
        check_cancelled()
    if options.csv_file:
        generated.append(write_csv(result, output_directory / f"{prefix}_逐帧遥测.csv"))
        check_cancelled()
    if options.xlsx_file:
        generated.append(write_xlsx(result, output_directory / f"{prefix}_逐帧遥测.xlsx"))
        check_cancelled()
    if options.camera_json:
        generated.append(_atomic_json(
            output_directory / f"{prefix}_相机内参与镜头信息.json",
            build_camera_lens_profile(result, source_video=source_video),
            compact=True,
        ))
        check_cancelled()
    if options.ar_audit_json:
        generated.append(_atomic_json(
            output_directory / f"{prefix}_AR投影数据准备.json",
            build_ar_preparation_document(
                result,
                source_video=source_video,
                projection_options=projection_options,
            ),
            compact=True,
        ))
        check_cancelled()
    if projection_options is not None and bool(getattr(projection_options, "enabled", False)):
        from projection_reconstruction import run_projection_reconstruction

        outcome = run_projection_reconstruction(
            video_path,
            source_video=source_video,
            samples=_projection_sample_records(result),
            output_directory=output_directory,
            options=projection_options,
            progress=progress_callback,
            cancel_event=cancel_event,
        )
        generated.append(outcome.report_path)
        if outcome.mesh_path is not None:
            generated.append(outcome.mesh_path)
        if outcome.status == "cancelled":
            raise ExtractionCancelled(
                outcome.message or "用户已取消三维重建",
                result=result,
                generated=generated,
            )
        projection_success_committed = outcome.status == "success"
        if not projection_success_committed:
            check_cancelled()
    report_path = output_directory / f"{prefix}_提取报告.txt"
    report = build_report(result, [*generated, report_path], timezone_label)
    generated.append(_atomic_bytes(report_path, report.encode("utf-8-sig")))
    if not projection_success_committed:
        check_cancelled()
    return result, generated


def _find_sidecar(video: Path) -> Path | None:
    if not video.parent.is_dir():
        return None
    candidates = [path for path in video.parent.iterdir() if path.is_file() and path.suffix.casefold() == ".srt"]
    exact = [path for path in candidates if path.stem.casefold() == video.stem.casefold()]
    # Prefix matching can silently bind telemetry from another flight.  Only a
    # unique, strictly same-stem sidecar is safe for automatic selection.
    return exact[0] if len(exact) == 1 else None


def _append_launcher_log(message: str) -> None:
    """Best-effort UTF-8 diagnostics shared by the CMD launcher and Python GUI."""

    try:
        timestamp = datetime.now().astimezone().isoformat(timespec="seconds")
        with LAUNCHER_LOG_PATH.open("a", encoding="utf-8", newline="\n") as handle:
            handle.write(f"\n[{timestamp}] {message.rstrip()}\n")
    except BaseException:
        # Logging must never hide the original failure.
        pass


def _show_fatal_gui_error(details: str) -> None:
    _append_launcher_log(f"未处理异常\n{details}")
    try:
        import tkinter as tk
        from tkinter import messagebox

        dialog_root = tk.Tk()
        dialog_root.withdraw()
        messagebox.showerror(
            "DJI 提取器启动失败",
            "程序发生未处理错误。完整 traceback 已写入：\n"
            f"{LAUNCHER_LOG_PATH}\n\n"
            f"错误摘要：{details.strip().splitlines()[-1] if details.strip() else '未知错误'}",
            parent=dialog_root,
        )
        dialog_root.destroy()
    except BaseException:
        pass


def launch_gui(*, smoke_test: bool = False) -> int:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
    from projection_reconstruction import ProjectionOptions, find_colmap_executable, projection_capabilities

    root = tk.Tk()
    root.title(f"DJI 航拍视频 AR 数据提取器 {APP_VERSION}")
    root.geometry("900x760")
    root.minsize(820, 680)

    def report_callback_exception(exception_type: type[BaseException], exception: BaseException, tb: Any) -> None:
        details = "".join(traceback.format_exception(exception_type, exception, tb))
        _append_launcher_log(f"Tk 回调未处理异常\n{details}")
        try:
            messagebox.showerror(
                "程序错误",
                "界面发生未处理错误。完整 traceback 已写入：\n"
                f"{LAUNCHER_LOG_PATH}\n\n{exception}",
                parent=root,
            )
        except BaseException:
            pass

    root.report_callback_exception = report_callback_exception

    video_var = tk.StringVar()
    srt_var = tk.StringVar()
    output_var = tk.StringVar()
    local_offset = datetime.now().astimezone().strftime("%z")
    timezone_var = tk.StringVar(value=f"{local_offset[:3]}:{local_offset[3:]}" if local_offset else "local")
    raw_var = tk.BooleanVar(value=True)
    cesium_var = tk.BooleanVar(value=True)
    csv_var = tk.BooleanVar(value=True)
    xlsx_var = tk.BooleanVar(value=True)
    camera_var = tk.BooleanVar(value=True)
    ar_audit_var = tk.BooleanVar(value=True)
    projection_var = tk.BooleanVar(value=False)
    detected_colmap = find_colmap_executable(None)
    colmap_var = tk.StringVar(value=str(detected_colmap) if detected_colmap else "colmap")
    height_mode_labels = {
        "未确认（不生成三维网格）": "unconfirmed",
        "DJI 绝对/GPS 高度就是 WGS84 椭球高": "absolute_is_wgs84_ellipsoid",
        "DJI 正高 + 大地水准面修正值": "orthometric_plus_geoid_offset",
        "DJI 相对高度 + 起飞点椭球高": "relative_plus_takeoff_ellipsoid",
    }
    height_mode_var = tk.StringVar(value=next(iter(height_mode_labels)))
    height_value_var = tk.StringVar(value="0.0")
    mesh_fps_var = tk.StringVar(value="12")
    max_dimension_var = tk.StringVar(value="1280")
    events: queue.Queue[tuple[str, Any]] = queue.Queue()
    cancel_event = threading.Event()
    worker_thread: threading.Thread | None = None
    pending_close = False

    frame = ttk.Frame(root, padding=16)
    frame.pack(fill="both", expand=True)
    frame.columnconfigure(1, weight=1)

    def choose_video() -> None:
        selected = filedialog.askopenfilename(
            title="选择 DJI 航拍视频",
            filetypes=[("航拍视频", "*.mp4 *.MP4 *.mov *.MOV *.m4v *.M4V"), ("全部文件", "*.*")],
        )
        if not selected:
            return
        video = Path(selected)
        # Never retain an SRT that belonged to the previously selected video.
        srt_var.set("")
        video_var.set(str(video))
        output_var.set(str(video.parent / f"{video.stem}_提取结果"))
        sidecar = _find_sidecar(video)
        if sidecar:
            srt_var.set(str(sidecar))

    def choose_srt() -> None:
        selected = filedialog.askopenfilename(title="选择 DJI SRT", filetypes=[("SRT字幕", "*.srt *.SRT"), ("全部文件", "*.*")])
        if selected:
            srt_var.set(selected)

    def choose_output() -> None:
        selected = filedialog.askdirectory(title="选择输出目录")
        if selected:
            output_var.set(selected)

    rows = [
        ("DJI 视频", video_var, choose_video, "选择视频"),
        ("DJI SRT（可留空）", srt_var, choose_srt, "选择 SRT"),
        ("输出目录", output_var, choose_output, "选择目录"),
    ]
    action_widgets: list[tk.Widget] = []
    for row_index, (label, variable, command, button_text) in enumerate(rows):
        ttk.Label(frame, text=label).grid(row=row_index, column=0, sticky="w", padx=(0, 10), pady=6)
        entry = ttk.Entry(frame, textvariable=variable)
        entry.grid(row=row_index, column=1, sticky="ew", pady=6)
        button = ttk.Button(frame, text=button_text, command=command)
        button.grid(row=row_index, column=2, padx=(10, 0), pady=6)
        action_widgets.extend([entry, button])

    ttk.Label(frame, text="SRT墙钟时区").grid(row=3, column=0, sticky="w", padx=(0, 10), pady=6)
    timezone_box = ttk.Combobox(frame, textvariable=timezone_var, values=("+08:00", "+00:00", "local"), width=16)
    timezone_box.grid(row=3, column=1, sticky="w", pady=6)
    action_widgets.append(timezone_box)

    options_frame = ttk.LabelFrame(frame, text="输出", padding=8)
    options_frame.grid(row=4, column=0, columnspan=3, sticky="ew", pady=(10, 8))
    for index, (text, variable) in enumerate((
        ("全部原始元数据 JSON", raw_var),
        ("Cesium 轨迹 JSON", cesium_var),
        ("逐帧 CSV", csv_var),
        ("逐帧 Excel", xlsx_var),
        ("相机/镜头 JSON", camera_var),
        ("AR 数据准备审计", ar_audit_var),
    )):
        widget = ttk.Checkbutton(options_frame, text=text, variable=variable)
        widget.grid(row=index // 4, column=index % 4, padx=8, pady=2, sticky="w")
        action_widgets.append(widget)

    reconstruction_frame = ttk.LabelFrame(frame, text="真实表面 AR 网格（可选，耗时）", padding=8)
    reconstruction_frame.grid(row=5, column=0, columnspan=3, sticky="ew", pady=(2, 8))
    reconstruction_frame.columnconfigure(1, weight=1)
    projection_toggle = ttk.Checkbutton(
        reconstruction_frame,
        text="运行 SfM/MVS，提取楼顶、立面和地面的逐帧三维可见表面",
        variable=projection_var,
    )
    projection_toggle.grid(row=0, column=0, columnspan=4, sticky="w", pady=(0, 6))
    action_widgets.append(projection_toggle)

    def choose_colmap() -> None:
        selected = filedialog.askopenfilename(
            title="选择 COLMAP.exe 或 COLMAP.bat",
            filetypes=[("COLMAP", "*.exe *.bat *.cmd"), ("全部文件", "*.*")],
        )
        if selected:
            colmap_var.set(selected)

    ttk.Label(reconstruction_frame, text="COLMAP").grid(row=1, column=0, sticky="w", padx=(0, 8), pady=3)
    colmap_entry = ttk.Entry(reconstruction_frame, textvariable=colmap_var)
    colmap_entry.grid(row=1, column=1, sticky="ew", pady=3)
    colmap_button = ttk.Button(reconstruction_frame, text="选择", command=choose_colmap)
    colmap_button.grid(row=1, column=2, padx=(8, 0), pady=3)
    action_widgets.extend([colmap_entry, colmap_button])

    ttk.Label(reconstruction_frame, text="高程基准").grid(row=2, column=0, sticky="w", padx=(0, 8), pady=3)
    height_mode_box = ttk.Combobox(
        reconstruction_frame,
        textvariable=height_mode_var,
        values=tuple(height_mode_labels),
        state="readonly",
    )
    height_mode_box.grid(row=2, column=1, columnspan=2, sticky="ew", pady=3)
    action_widgets.append(height_mode_box)

    ttk.Label(reconstruction_frame, text="修正值/起飞点椭球高(m)").grid(row=3, column=0, sticky="w", padx=(0, 8), pady=3)
    height_entry = ttk.Entry(reconstruction_frame, textvariable=height_value_var, width=14)
    height_entry.grid(row=3, column=1, sticky="w", pady=3)
    ttk.Label(reconstruction_frame, text="网格 FPS").grid(row=3, column=1, sticky="e", padx=(0, 76), pady=3)
    mesh_fps_entry = ttk.Entry(reconstruction_frame, textvariable=mesh_fps_var, width=7)
    mesh_fps_entry.grid(row=3, column=1, sticky="e", pady=3)
    ttk.Label(reconstruction_frame, text="重建最长边").grid(row=3, column=2, sticky="e", padx=(8, 72), pady=3)
    max_dimension_entry = ttk.Entry(reconstruction_frame, textvariable=max_dimension_var, width=7)
    max_dimension_entry.grid(row=3, column=2, sticky="e", pady=3)
    action_widgets.extend([height_entry, mesh_fps_entry, max_dimension_entry])

    capability_var = tk.StringVar()

    def refresh_capabilities() -> None:
        options = ProjectionOptions(
            enabled=projection_var.get(),
            colmap_bin=colmap_var.get().strip() or "colmap",
            height_mode=height_mode_labels.get(height_mode_var.get(), "unconfirmed"),
        )
        capabilities = projection_capabilities(options)
        capability_var.set(
            "重建环境已就绪" if capabilities["ready"]
            else "缺少：" + "、".join(capabilities["missing"])
        )

    capability_button = ttk.Button(reconstruction_frame, text="检查环境", command=refresh_capabilities)
    capability_button.grid(row=1, column=3, padx=(8, 0), pady=3)
    ttk.Label(reconstruction_frame, textvariable=capability_var, foreground="#8A4B08").grid(
        row=4, column=0, columnspan=4, sticky="w", pady=(4, 0)
    )
    action_widgets.append(capability_button)
    refresh_capabilities()

    hint = (
        "基础模式直接提取 DJI 遥测、估算内参并输出缺项审计。真实表面模式会在后台运行 COLMAP，"
        "只有全部帧、镜头标定、首表面深度和 EPSG:4979 地理配准通过校验后才发布网格；失败不会伪造地面点。"
    )
    ttk.Label(frame, text=hint, foreground="#555555", wraplength=850).grid(row=6, column=0, columnspan=3, sticky="w", pady=6)

    progress = ttk.Progressbar(frame, mode="indeterminate")
    progress.grid(row=7, column=0, columnspan=3, sticky="ew", pady=(8, 4))
    log = tk.Text(frame, height=11, wrap="word", state="disabled")
    log.grid(row=8, column=0, columnspan=3, sticky="nsew", pady=6)
    frame.rowconfigure(8, weight=1)

    def append_log(text: str) -> None:
        log.configure(state="normal")
        log.insert("end", text.rstrip() + "\n")
        try:
            line_count = int(log.index("end-1c").split(".")[0])
            if line_count > 2000:
                log.delete("1.0", f"{line_count - 1800}.0")
        except (ValueError, tk.TclError):
            pass
        log.see("end")
        log.configure(state="disabled")

    def set_busy(busy: bool) -> None:
        for widget in action_widgets:
            try:
                if busy:
                    widget.configure(state="disabled")
                else:
                    widget.configure(state="readonly" if widget is height_mode_box else "normal")
            except tk.TclError:
                pass
        extract_button.configure(state="disabled" if busy else "normal")
        cancel_button.configure(state="normal" if busy else "disabled")
        if busy:
            progress.start(12)
        else:
            progress.stop()

    def worker(
        video: Path,
        srt: Path | None,
        output: Path,
        timezone_text: str,
        export_options: ExportOptions,
        ar_options: ProjectionOptions,
    ) -> None:
        def on_progress(stage: str, message: str, fraction: float | None) -> None:
            events.put(("progress", (stage, message, fraction)))

        try:
            result, generated = run_extraction(
                video,
                srt_path=srt,
                output_directory=output,
                source_timezone_text=timezone_text,
                options=export_options,
                projection_options=ar_options,
                progress_callback=on_progress,
                cancel_event=cancel_event,
            )
            events.put(("success", (result, generated)))
        except ExtractionCancelled as exc:
            events.put(("cancelled", exc))
        except BaseException as exc:
            _append_launcher_log(f"提取任务失败\n{traceback.format_exc()}")
            events.put(("error", exc))

    def start() -> None:
        nonlocal worker_thread
        video_text = video_var.get().strip()
        srt_text = srt_var.get().strip()
        output_text = output_var.get().strip()
        if not video_text or not Path(video_text).is_file():
            messagebox.showerror("无法开始", "请先选择有效的 DJI MP4/MOV 视频")
            return
        if srt_text and not Path(srt_text).is_file():
            messagebox.showerror("无法开始", "指定的 SRT 文件不存在")
            return
        if not output_text:
            messagebox.showerror("无法开始", "请选择输出目录")
            return
        try:
            height_value = float(height_value_var.get().strip() or "0")
            mesh_fps = float(mesh_fps_var.get().strip())
            max_dimension = int(max_dimension_var.get().strip())
            ar_options = ProjectionOptions(
                enabled=projection_var.get(),
                colmap_bin=colmap_var.get().strip() or "colmap",
                mesh_fps=mesh_fps,
                max_image_dimension=max_dimension,
                height_mode=height_mode_labels.get(height_mode_var.get(), "unconfirmed"),
                height_offset_m=height_value,
            )
            ar_options.validate()
        except (ValueError, TypeError) as exc:
            messagebox.showerror("AR 参数无效", str(exc))
            return
        if ar_options.enabled and ar_options.height_mode == "unconfirmed":
            messagebox.showerror(
                "无法生成 EPSG:4979 网格",
                "请先确认 DJI 高度的垂直基准。未知基准不能冒充 WGS84 椭球高。",
            )
            return
        worker_arguments = (
            Path(video_text),
            Path(srt_text) if srt_text else None,
            Path(output_text),
            timezone_var.get(),
            ExportOptions(
                raw_var.get(), cesium_var.get(), csv_var.get(), xlsx_var.get(),
                camera_var.get(), ar_audit_var.get(),
            ),
            ar_options,
        )
        cancel_event.clear()
        append_log("开始提取 DJI 遥测和 AR 所需相机信息……")
        if ar_options.enabled:
            append_log("已启用真实表面 SfM/MVS；任务可能持续较长时间，可随时点击取消。")
        set_busy(True)
        worker_thread = threading.Thread(target=worker, args=worker_arguments, name="dji-extractor", daemon=False)
        worker_thread.start()

    def poll() -> None:
        nonlocal pending_close
        terminal = False
        for _ in range(40):
            try:
                kind, payload = events.get_nowait()
            except queue.Empty:
                break
            if kind == "progress":
                stage, message, _fraction = payload
                append_log(f"[{stage}] {message}")
                continue
            terminal = True
            set_busy(False)
            if kind == "success":
                result, generated = payload
                sample_count = len(_primary_samples(result))
                ar_readiness = _ar_projection_contract(
                    _projection_audits(_primary_samples(result), result.video.width, result.video.height)
                )["readiness"]
                ar_message = (
                    "姿态估算输入完整"
                    if ar_readiness["ready"]
                    else (
                        "姿态估算输入不完整，完整输入帧 "
                        f"{ar_readiness['ready_point_count']}/{ar_readiness['source_sample_count']}"
                    )
                )
                cesium_path = next((path for path in generated if "Cesium轨迹" in path.name), None)
                mesh_path = next((path for path in generated if "三维投影网格_请选择此文件" in path.name), None)
                quality_path = next((path for path in generated if "AR投影质量报告" in path.name), None)
                reconstruction_message = "未运行真实表面重建"
                reconstruction_status = "not_requested"
                if quality_path and quality_path.is_file():
                    try:
                        quality = json.loads(quality_path.read_text(encoding="utf-8"))
                        reconstruction_status = str(quality.get("status") or "unknown")
                        reconstruction_message = str(quality.get("message") or reconstruction_status)
                    except (OSError, json.JSONDecodeError) as exc:
                        reconstruction_status = "failed"
                        reconstruction_message = f"无法读取 AR 质量报告：{exc}"
                append_log(f"完成：提取 {sample_count} 条逐帧遥测；{ar_message}")
                append_log(f"真实表面重建：{reconstruction_message}")
                for path in generated:
                    append_log(str(path))
                for warning in result.warnings:
                    append_log(f"警告：{warning}")
                if pending_close:
                    root.destroy()
                    return
                message = (
                    f"已提取 {sample_count} 条记录\n\n"
                    f"轨迹 JSON：\n{cesium_path.name if cesium_path else '本次未生成'}\n\n"
                    f"真实表面网格：\n{mesh_path.name if mesh_path else '未生成'}\n"
                    f"{reconstruction_message}\n\n"
                    "Cesium 的三维网格输入只选择“*_三维投影网格_请选择此文件.json”。"
                )
                if reconstruction_status in {"failed", "cancelled"}:
                    messagebox.showwarning("基础数据已完成，三维网格未生成", message)
                else:
                    messagebox.showinfo("提取完成", message)
            elif kind == "cancelled":
                append_log(f"已取消：{payload}")
                for path in payload.generated:
                    append_log(str(path))
                if pending_close:
                    root.destroy()
                    return
                messagebox.showinfo("任务已取消", str(payload))
            else:
                append_log(f"失败：{payload}")
                if pending_close:
                    root.destroy()
                    return
                messagebox.showerror("提取失败", str(payload))
        if pending_close and not terminal and (worker_thread is None or not worker_thread.is_alive()):
            root.destroy()
            return
        root.after(100, poll)

    def cancel() -> None:
        if worker_thread is not None and worker_thread.is_alive():
            cancel_event.set()
            cancel_button.configure(state="disabled")
            append_log("正在取消任务；当前外部进程退出后会保留质量报告和诊断工作区……")

    def close_window() -> None:
        nonlocal pending_close
        if worker_thread is not None and worker_thread.is_alive():
            if not messagebox.askyesno("任务仍在运行", "要取消当前任务并关闭窗口吗？"):
                return
            pending_close = True
            cancel()
            return
        root.destroy()

    extract_button = ttk.Button(frame, text="开始提取", command=start)
    extract_button.grid(row=9, column=1, sticky="e", padx=(0, 80), pady=(10, 0))
    cancel_button = ttk.Button(frame, text="取消", command=cancel, state="disabled")
    cancel_button.grid(row=9, column=1, sticky="e", pady=(10, 0))
    root.protocol("WM_DELETE_WINDOW", close_window)
    root.after(100, poll)
    if smoke_test:
        root.withdraw()
        root.after(250, root.destroy)
    root.mainloop()
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="DJI 航拍视频信息提取器")
    parser.add_argument("video", type=Path, nargs="?", help="DJI MP4/MOV；省略时打开图形界面")
    parser.add_argument("-s", "--srt", type=Path, help="配套 DJI SRT；省略时自动查找同名文件")
    parser.add_argument("-o", "--output-dir", type=Path, help="输出目录，默认在视频旁创建 *_提取结果")
    parser.add_argument("--timezone", default="local", help="SRT墙钟时区：local、UTC 或 +08:00")
    parser.add_argument("--ffprobe", default="ffprobe", help="可选 ffprobe 路径")
    parser.add_argument("--ffmpeg", default="ffmpeg", help="可选 ffmpeg 路径")
    parser.add_argument("--no-raw-json", action="store_true")
    parser.add_argument("--no-cesium-json", action="store_true")
    parser.add_argument("--no-csv", action="store_true")
    parser.add_argument("--no-xlsx", action="store_true")
    parser.add_argument("--no-camera-json", action="store_true")
    parser.add_argument("--no-ar-audit-json", action="store_true")
    parser.add_argument("--ar-mesh", action="store_true", help="运行 COLMAP SfM/MVS 并生成逐帧真实表面三维网格")
    parser.add_argument("--colmap", default="colmap", help="COLMAP.exe 或 COLMAP.bat 路径")
    parser.add_argument(
        "--height-mode",
        choices=(
            "unconfirmed",
            "absolute_is_wgs84_ellipsoid",
            "orthometric_plus_geoid_offset",
            "relative_plus_takeoff_ellipsoid",
        ),
        default="unconfirmed",
        help="输出 EPSG:4979 前必须明确 DJI 高程基准",
    )
    parser.add_argument("--height-offset-m", type=float, default=0.0, help="大地水准面修正值或起飞点椭球高")
    parser.add_argument("--mesh-fps", type=float, default=12.0, help="精确网格采样率，10–30")
    parser.add_argument("--reconstruction-max-dimension", type=int, default=1280)
    parser.add_argument("--gui", action="store_true", help="打开图形界面")
    parser.add_argument("--gui-smoke-test", action="store_true", help=argparse.SUPPRESS)
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    if args.gui_smoke_test:
        return launch_gui(smoke_test=True)
    if args.gui or args.video is None:
        return launch_gui()
    video = args.video.expanduser().resolve()
    output = args.output_dir or video.parent / f"{video.stem}_提取结果"
    try:
        from projection_reconstruction import ProjectionOptions

        projection_options = ProjectionOptions(
            enabled=args.ar_mesh,
            colmap_bin=args.colmap,
            mesh_fps=args.mesh_fps,
            max_image_dimension=args.reconstruction_max_dimension,
            height_mode=args.height_mode,
            height_offset_m=args.height_offset_m,
        )
        projection_options.validate()
        result, generated = run_extraction(
            video,
            srt_path=args.srt,
            output_directory=output,
            source_timezone_text=args.timezone,
            options=ExportOptions(
                raw_json=not args.no_raw_json,
                cesium_json=not args.no_cesium_json,
                csv_file=not args.no_csv,
                xlsx_file=not args.no_xlsx,
                camera_json=not args.no_camera_json,
                ar_audit_json=not args.no_ar_audit_json,
            ),
            ffprobe_bin=args.ffprobe,
            ffmpeg_bin=args.ffmpeg,
            projection_options=projection_options,
            progress_callback=lambda stage, message, _fraction: print(f"[{stage}] {message}"),
        )
    except (ExtractionError, OSError, RuntimeError, ValueError) as exc:
        print(f"提取失败：{exc}", file=sys.stderr)
        return 2
    print(f"已提取 {len(_primary_samples(result))} 条逐帧遥测：")
    for path in generated:
        print(f"  {path}")
    for warning in result.warnings:
        print(f"警告：{warning}")
    if args.ar_mesh:
        quality_path = next((path for path in generated if "AR投影质量报告" in path.name), None)
        reconstruction_status = "missing_report"
        reconstruction_message = "未生成 AR 投影质量报告"
        if quality_path is not None and quality_path.is_file():
            try:
                quality = json.loads(quality_path.read_text(encoding="utf-8"))
                reconstruction_status = str(quality.get("status") or "unknown")
                reconstruction_message = str(quality.get("message") or reconstruction_status)
            except (OSError, json.JSONDecodeError) as exc:
                reconstruction_message = f"无法读取 AR 投影质量报告：{exc}"
        if reconstruction_status != "success":
            print(f"真实表面 AR 网格未生成：{reconstruction_message}", file=sys.stderr)
            return 3
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except SystemExit:
        raise
    except BaseException:
        fatal_details = traceback.format_exc()
        print(fatal_details, file=sys.stderr)
        _show_fatal_gui_error(fatal_details)
        raise SystemExit(1)
