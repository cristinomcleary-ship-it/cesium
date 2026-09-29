from __future__ import annotations

import csv
import json
import os
import tempfile
import threading
import unittest
from dataclasses import replace
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

import dji_video_extractor as wrapper
from backend_v1.dji_telemetry_extractor import (
    ExtractionCancelled as BackendExtractionCancelled,
    ExtractionResult,
    TelemetrySample,
    VideoMetadata,
)


def fixture_result() -> ExtractionResult:
    video = VideoMetadata(
        backend="fixture",
        width=3840,
        height=2160,
        frame_rate_fps=30_000 / 1_001,
        frame_count=2,
        duration_ms=67,
        video_codec="hevc",
    )
    common = {
        "source_kind": "sidecar_srt",
        "source_name": "DJI_0001.SRT",
        "position_crs": "EPSG:4326",
        "absolute_altitude_m": 99.2,
        "relative_altitude_m": 70.1,
        "gimbal_yaw_deg": -102.9,
        "gimbal_pitch_deg": -30.0,
        "gimbal_roll_deg": 0.0,
        "focal_length_35mm_equivalent_mm": 24.0,
        "digital_zoom_ratio": 1.0,
        "iso": 280,
    }
    return ExtractionResult(
        source_file="DJI_0001.MP4",
        video=video,
        telemetry=(
            TelemetrySample(
                **common,
                cue_index=1,
                video_sample_index=0,
                frame_index=0,
                start_time_ms=0,
                end_time_ms=33,
                recorded_at="2026-08-19T16:20:16.562",
                latitude=30.836687,
                longitude=120.890825,
                extra_fields={"dehaze_level": "0"},
            ),
            TelemetrySample(
                **common,
                cue_index=2,
                video_sample_index=1,
                # Real DJI FrameCnt may skip values; it is not a decoder index.
                frame_index=3,
                start_time_ms=33,
                end_time_ms=67,
                recorded_at="2026-08-19T16:20:16.613",
                latitude=30.836688,
                longitude=120.890826,
                extra_fields={"dehaze_level": "1"},
            ),
        ),
    )


class CesiumTrajectoryContractTests(unittest.TestCase):
    def test_trajectory_uses_contiguous_samples_and_keeps_original_framecnt(self) -> None:
        document = wrapper.build_cesium_trajectory(
            fixture_result(),
            source_timezone=timezone(timedelta(hours=8)),
            timezone_label="+08:00",
        )
        points = document["trajectory_points"]
        self.assertEqual([point["video_sample_index"] for point in points], [0, 1])
        self.assertEqual([point["frame_index"] for point in points], [0, 3])
        self.assertEqual([point["video_time_ms"] for point in points], [0, 33])

    def test_pts_timeline_is_anchored_once_to_first_wall_clock_sample(self) -> None:
        source_timezone = timezone(timedelta(hours=8))
        document = wrapper.build_cesium_trajectory(
            fixture_result(),
            source_timezone=source_timezone,
            timezone_label="+08:00",
        )
        first, second = document["trajectory_points"]
        first_wall_clock_ms = round(
            datetime(2026, 8, 19, 16, 20, 16, 562000, tzinfo=source_timezone).timestamp() * 1000
        )
        second_wall_clock_ms = round(
            datetime(2026, 8, 19, 16, 20, 16, 613000, tzinfo=source_timezone).timestamp() * 1000
        )
        self.assertEqual(first["source_timestamp_ms"], first_wall_clock_ms)
        self.assertEqual(second["source_timestamp_ms"], first_wall_clock_ms + 33)
        self.assertNotEqual(second["source_timestamp_ms"], second_wall_clock_ms)
        self.assertEqual(second["video_frame_timestamp_ms"], first_wall_clock_ms + 33)
        self.assertEqual(second["video_frame_time_delta_ms"], 0)
        self.assertEqual(
            document["flight_info"]["timeline"],
            "exact SRT PTS anchored to first DJI wall-clock sample",
        )

    def test_missing_aircraft_attitude_is_not_filled_from_gimbal(self) -> None:
        document = wrapper.build_cesium_trajectory(
            fixture_result(),
            source_timezone=timezone.utc,
            timezone_label="+00:00",
        )
        self.assertEqual(document["flight_info"]["body_attitude_status"], "not_present_in_DJI_SRT")
        for point in document["trajectory_points"]:
            self.assertNotIn("attitude", point)
            self.assertEqual(point["gimbal_attitude"]["pitch"], -30.0)

    def test_equivalent_focal_length_yields_explicitly_estimated_intrinsics(self) -> None:
        document = wrapper.build_cesium_trajectory(
            fixture_result(),
            source_timezone=timezone.utc,
            timezone_label="+00:00",
        )
        camera_summary = document["flight_info"]["camera_model"]
        self.assertEqual(camera_summary["intrinsics_status"], "estimated_from_35mm_equivalent")
        self.assertEqual(camera_summary["lens_distortion_status"], "not_present_in_DJI_SRT")
        camera = document["trajectory_points"][0]["camera"]
        self.assertEqual(camera["focal_length_35mm_equivalent_mm"], 24.0)
        estimated = camera["estimated_intrinsics"]
        self.assertEqual(estimated["quality"], "estimated_not_calibrated")
        self.assertIn("35mm-equivalent", estimated["method"])
        self.assertEqual(estimated["fx_px"], 2560.0)
        self.assertNotIn("distortion_coefficients", estimated)

    def test_trajectory_does_not_claim_unmeasured_ground_corners(self) -> None:
        document = wrapper.build_cesium_trajectory(
            fixture_result(),
            source_timezone=timezone.utc,
            timezone_label="+00:00",
        )
        self.assertEqual(document["source_crs"], "EPSG:4326")
        self.assertIn("not_extracted", document["flight_info"]["ground_footprint_status"])
        wire = json.dumps(document, ensure_ascii=False)
        self.assertNotIn("ground_footprint_wgs84", wire)
        self.assertNotIn("top_left", wire)
        self.assertNotIn("top_right", wire)
        self.assertNotIn("bottom_left", wire)
        self.assertNotIn("bottom_right", wire)

    def test_ar_projection_contract_is_ready_only_with_complete_explicit_inputs(self) -> None:
        document = wrapper.build_cesium_trajectory(
            fixture_result(),
            source_timezone=timezone.utc,
            timezone_label="+00:00",
        )
        projection = document["flight_info"]["ar_projection"]
        readiness = projection["readiness"]
        self.assertEqual(projection["mode"], "dji_srt_pose_terrain_intersection")
        self.assertTrue(readiness["ready"])
        self.assertEqual(readiness["status"], "ready_estimated")
        self.assertEqual(readiness["source_sample_count"], 2)
        self.assertEqual(readiness["ready_point_count"], 2)
        self.assertTrue(projection["runtime_requirements"]["terrain_intersection"])
        self.assertFalse(projection["measured_outputs"]["ground_corners_available"])
        for point in document["trajectory_points"]:
            self.assertTrue(point["ar_projection_ready"])
            self.assertNotIn("ar_projection_missing_inputs", point)

    def test_ar_projection_contract_reports_partial_frames_and_exact_missing_inputs(self) -> None:
        result = fixture_result()
        samples = list(result.telemetry)
        samples[1] = replace(samples[1], gimbal_roll_deg=None)
        partial = replace(result, telemetry=tuple(samples))
        document = wrapper.build_cesium_trajectory(
            partial,
            source_timezone=timezone.utc,
            timezone_label="+00:00",
        )
        readiness = document["flight_info"]["ar_projection"]["readiness"]
        self.assertFalse(readiness["ready"])
        self.assertEqual(readiness["status"], "partial")
        self.assertEqual(readiness["ready_point_count"], 1)
        self.assertEqual(readiness["missing_input_counts"]["gimbal_yaw_pitch_roll"], 1)
        second = document["trajectory_points"][1]
        self.assertFalse(second["ar_projection_ready"])
        self.assertEqual(second["ar_projection_missing_inputs"], ["gimbal_yaw_pitch_roll"])

    def test_compatibility_timeline_fallbacks_do_not_claim_projection_readiness(self) -> None:
        result = fixture_result()
        samples = list(result.telemetry)
        samples[0] = replace(samples[0], video_sample_index=None, start_time_ms=None)
        incomplete = replace(result, telemetry=tuple(samples))
        document = wrapper.build_cesium_trajectory(
            incomplete,
            source_timezone=timezone.utc,
            timezone_label="+00:00",
        )
        first = document["trajectory_points"][0]
        self.assertEqual(first["video_sample_index"], 0)
        self.assertEqual(first["video_time_ms"], 0)
        self.assertFalse(first["ar_projection_ready"])
        self.assertEqual(
            first["ar_projection_missing_inputs"],
            ["video_sample_index", "video_time_ms"],
        )


class RunExtractionCancellationTests(unittest.TestCase):
    def test_cancelled_reconstruction_raises_normal_cancellation_with_report(self) -> None:
        from projection_reconstruction import ProjectionOptions, ProjectionOutcome

        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            video = directory / "DJI_0001.MP4"
            video.write_bytes(b"fixture-video")
            output = directory / "output"
            report = output / "projection-quality.json"
            cancel_event = threading.Event()

            def cancelled_reconstruction(*_args: object, **_kwargs: object) -> ProjectionOutcome:
                report.parent.mkdir(parents=True, exist_ok=True)
                report.write_text('{"status":"cancelled"}\n', encoding="utf-8")
                cancel_event.set()
                return ProjectionOutcome("cancelled", report, message="cancelled by test")

            with (
                patch.object(wrapper, "extract_dji_video", return_value=fixture_result()) as extract,
                patch(
                    "projection_reconstruction.run_projection_reconstruction",
                    side_effect=cancelled_reconstruction,
                ),
            ):
                with self.assertRaises(wrapper.ExtractionCancelled) as raised:
                    wrapper.run_extraction(
                        video,
                        srt_path=None,
                        output_directory=output,
                        source_timezone_text="UTC",
                        options=wrapper.ExportOptions(False, False, False, False, False, False),
                        projection_options=ProjectionOptions(
                            enabled=True,
                            height_mode="absolute_is_wgs84_ellipsoid",
                        ),
                        cancel_event=cancel_event,
                    )

            self.assertEqual(str(raised.exception), "cancelled by test")
            self.assertEqual(raised.exception.result, fixture_result())
            self.assertEqual(raised.exception.generated, (report,))
            self.assertIs(extract.call_args.kwargs["cancel_event"], cancel_event)

    def test_telemetry_cancellation_is_converted_to_the_gui_terminal_type(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            video = directory / "DJI_0001.MP4"
            video.write_bytes(b"fixture-video")
            cancel_event = threading.Event()

            with patch.object(
                wrapper,
                "extract_dji_video",
                side_effect=BackendExtractionCancelled("cancelled during ffprobe"),
            ) as extract:
                with self.assertRaises(wrapper.ExtractionCancelled) as raised:
                    wrapper.run_extraction(
                        video,
                        srt_path=None,
                        output_directory=directory / "output",
                        source_timezone_text="UTC",
                        options=wrapper.ExportOptions(False, False, False, False, False, False),
                        cancel_event=cancel_event,
                    )

            self.assertEqual(str(raised.exception), "cancelled during ffprobe")
            self.assertIsInstance(raised.exception.__cause__, BackendExtractionCancelled)
            self.assertIs(extract.call_args.kwargs["cancel_event"], cancel_event)

    def test_late_cancel_does_not_retroactively_reject_published_mesh(self) -> None:
        from projection_reconstruction import ProjectionOptions, ProjectionOutcome

        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            video = directory / "DJI_0001.MP4"
            video.write_bytes(b"fixture-video")
            output = directory / "output"
            report = output / "projection-quality.json"
            mesh = output / "projection-mesh.json"
            cancel_event = threading.Event()

            def successful_reconstruction(*_args: object, **_kwargs: object) -> ProjectionOutcome:
                report.parent.mkdir(parents=True, exist_ok=True)
                report.write_text('{"status":"success"}\n', encoding="utf-8")
                mesh.write_text('{"schema_version":"test"}\n', encoding="utf-8")
                cancel_event.set()
                return ProjectionOutcome("success", report, mesh, message="published")

            expected = fixture_result()
            with (
                patch.object(wrapper, "extract_dji_video", return_value=expected),
                patch(
                    "projection_reconstruction.run_projection_reconstruction",
                    side_effect=successful_reconstruction,
                ),
            ):
                result, generated = wrapper.run_extraction(
                    video,
                    srt_path=None,
                    output_directory=output,
                    source_timezone_text="UTC",
                    options=wrapper.ExportOptions(False, False, False, False, False, False),
                    projection_options=ProjectionOptions(
                        enabled=True,
                        height_mode="absolute_is_wgs84_ellipsoid",
                    ),
                    cancel_event=cancel_event,
                )

            self.assertEqual(result, expected)
            self.assertIn(report, generated)
            self.assertIn(mesh, generated)
            self.assertTrue(any(path.suffix == ".txt" and path.is_file() for path in generated))


class WrapperOutputTests(unittest.TestCase):
    def test_camera_and_ar_audit_outputs_are_explicit_about_estimates_and_surface_reconstruction(self) -> None:
        result = fixture_result()
        source_video = {
            "file_name": "DJI_0001.MP4",
            "size_bytes": 123456,
            "duration_ms": 67,
            "rotation_degrees": None,
        }
        camera = wrapper.build_camera_lens_profile(result, source_video=source_video)
        audit = wrapper.build_ar_preparation_document(result, source_video=source_video)

        self.assertEqual(camera["schema_version"], "dji-video-camera-lens/v2")
        self.assertEqual(camera["intrinsics"]["status"], "estimated_from_35mm_equivalent")
        self.assertTrue(camera["lens_distortion"]["must_not_assume_zero_for_exact_projection"])
        self.assertEqual(audit["schema_version"], "dji-video-ar-preparation/v2")
        self.assertEqual(audit["directly_extracted"]["exact_pts"], 2)
        self.assertFalse(audit["safety"]["flatten_buildings_to_ground"])
        self.assertEqual(audit["required_projection_mesh_contract"]["source_crs"], "EPSG:4979")

    def test_source_video_identity_carries_rotation_for_browser_space_preflight(self) -> None:
        result = fixture_result()
        rotated = replace(result, video=replace(result.video, rotation_deg=90.0))
        with tempfile.TemporaryDirectory() as temporary_directory:
            video = Path(temporary_directory) / "DJI_0001.MP4"
            video.write_bytes(b"fixture")
            identity = wrapper._source_video_identity(video, rotated)
        self.assertEqual(identity["rotation_degrees"], 90.0)
        self.assertEqual(len(identity["content_fingerprint_sha256"]), 64)
        self.assertIn("first/middle/last", identity["content_fingerprint_scheme"])

    def test_report_explains_estimated_ar_readiness_and_keeps_corner_input_empty(self) -> None:
        report = wrapper.build_report(
            fixture_result(),
            [Path("DJI_Cesium轨迹_请选择此文件.json")],
            "+08:00",
        )
        self.assertIn("AR 投影就绪状态: 已就绪", report)
        self.assertIn("SRT 姿态 + Cesium 地形求交，估算", report)
        self.assertIn("逐帧 WGS84 地面四角：本次未生成，请留空", report)

    def test_csv_keeps_sample_index_framecnt_and_raw_fields(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            destination = Path(temporary_directory) / "telemetry.csv"
            destination.write_text("old partial output", encoding="utf-8")
            original_replace = os.replace
            with (
                patch.object(wrapper.os, "replace", wraps=original_replace) as replace,
                patch.object(
                    wrapper,
                    "_csv_rows",
                    side_effect=AssertionError("streaming CSV must not materialize all rows"),
                ),
            ):
                written = wrapper.write_csv(fixture_result(), destination)

            replace.assert_called_once()
            source, target = replace.call_args.args
            self.assertTrue(str(source).endswith(".csv.tmp"))
            self.assertEqual(Path(target), destination.resolve())
            self.assertEqual(written, destination.resolve())
            with destination.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual([row["video_sample_index"] for row in rows], ["0", "1"])
            self.assertEqual([row["frame_index"] for row in rows], ["0", "3"])
            self.assertEqual([row["raw_dehaze_level"] for row in rows], ["0", "1"])
            self.assertEqual(list(Path(temporary_directory).glob("*.tmp")), [])

    def test_csv_rows_compatibility_helper_still_returns_materialized_rows(self) -> None:
        headers, rows = wrapper._csv_rows(fixture_result())

        self.assertIsInstance(headers, list)
        self.assertIsInstance(rows, list)
        self.assertEqual([row["video_sample_index"] for row in rows], [0, 1])
        self.assertIn("raw_dehaze_level", headers)

    def test_xlsx_streams_rows_and_preserves_header_sheet_features(self) -> None:
        from openpyxl import Workbook, load_workbook
        from openpyxl.utils import get_column_letter

        with tempfile.TemporaryDirectory() as temporary_directory:
            destination = Path(temporary_directory) / "telemetry.xlsx"
            with (
                patch("openpyxl.Workbook", wraps=Workbook) as workbook_constructor,
                patch.object(
                    wrapper,
                    "_csv_rows",
                    side_effect=AssertionError("streaming XLSX must not materialize all rows"),
                ),
            ):
                written = wrapper.write_xlsx(fixture_result(), destination)

            workbook_constructor.assert_called_once_with(write_only=True)
            self.assertEqual(written, destination.resolve())
            headers, rows = wrapper._csv_rows(fixture_result())
            workbook = load_workbook(destination, read_only=False)
            try:
                sheet = workbook.active
                self.assertEqual(sheet.freeze_panes, "A2")
                self.assertEqual(
                    sheet.auto_filter.ref,
                    f"A1:{get_column_letter(len(headers))}{len(rows) + 1}",
                )
                self.assertEqual(
                    [sheet.cell(1, index).value for index in range(1, len(headers) + 1)],
                    headers,
                )
                self.assertTrue(sheet["A1"].font.bold)
                self.assertTrue(str(sheet["A1"].font.color.rgb).endswith("FFFFFF"))
                self.assertTrue(str(sheet["A1"].fill.fgColor.rgb).endswith("1F4E78"))
                self.assertEqual(sheet.cell(2, headers.index("video_sample_index") + 1).value, 0)
                self.assertEqual(sheet.cell(3, headers.index("video_sample_index") + 1).value, 1)
                self.assertGreaterEqual(sheet.column_dimensions["A"].width, 12)
            finally:
                workbook.close()

    def test_launcher_is_ascii_and_keeps_a_diagnostic_console_open(self) -> None:
        entry_path = wrapper.SERVICE_ROOT / "提取大疆航拍视频信息.cmd"
        worker_path = wrapper.SERVICE_ROOT / "DJI_Extractor_Launcher.cmd"
        entry_bytes = entry_path.read_bytes()
        worker_bytes = worker_path.read_bytes()
        entry = entry_bytes.decode("ascii")
        worker = worker_bytes.decode("ascii")
        folded = worker.casefold()

        self.assertTrue(all(byte < 128 for byte in entry_bytes + worker_bytes))
        for payload in (entry_bytes, worker_bytes):
            self.assertEqual(payload.count(b"\n"), payload.count(b"\r\n"))
        self.assertIn('"%ComSpec%" /d /k call "%~dp0DJI_Extractor_Launcher.cmd"', entry)
        self.assertIn('set "DJI_LAUNCHER_HOSTED=1"', entry)
        self.assertIn("--self-test", entry)
        self.assertNotIn("pythonw", folded)
        self.assertIn("dji_extractor_launcher.log", worker)
        self.assertIn("LOCAL_TCL_ROOT", worker)
        self.assertIn("TCL_LIBRARY", worker)
        self.assertIn("TK_LIBRARY", worker)
        self.assertIn("root=tk.Tk()", worker)
        self.assertIn("import dji_video_extractor as app", worker)
        self.assertIn("SELFTEST_OK", worker)
        self.assertIn("pause >nul", folded)
        self.assertIn('"%PYTHON_EXE%" "%EXTRACTOR_SCRIPT%" --gui', worker)

    def test_python_diagnostic_log_is_utf8(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            destination = Path(temporary_directory) / "launcher.log"
            with patch.object(wrapper, "LAUNCHER_LOG_PATH", destination):
                wrapper._append_launcher_log("测试 traceback：镜头参数缺失")
            text = destination.read_text(encoding="utf-8")
            self.assertIn("测试 traceback：镜头参数缺失", text)

    def test_atomic_json_failure_preserves_previous_file_and_cleans_temp(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            destination = Path(temporary_directory) / "result.json"
            destination.write_text('{"state":"previous"}\n', encoding="utf-8")
            with patch.object(wrapper.os, "replace", side_effect=OSError("simulated replace failure")):
                with self.assertRaises(OSError):
                    wrapper._atomic_json(destination, {"state": "new"})

            self.assertEqual(
                json.loads(destination.read_text(encoding="utf-8")),
                {"state": "previous"},
            )
            self.assertEqual(list(Path(temporary_directory).glob("*.tmp")), [])


if __name__ == "__main__":
    unittest.main()
